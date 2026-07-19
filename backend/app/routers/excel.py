"""Excel 匯入/匯出。

匯入採「我們定義範本、公司照格式填」的策略（資料清洗成本轉嫁給提供方）：
- 必填欄極少：活動名稱/開始日/藝人/品項/商品名稱/售價/製作量 + 實售量或剩餘量擇一
- 全部驗證通過才寫入（all-or-nothing），先 dry_run 回報逐列錯誤
- 同一「活動×藝人×商品名」的多列視為同商品的多個規格
- 匯入的銷售異動 source=excel，與手動輸入區分
"""

from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app import schemas
from app.deps import get_company_id, get_db
from app.models import (
    Artist,
    ArtistMetric,
    Event,
    Expense,
    InventoryMovement,
    ItemType,
    MovementType,
    Product,
    ProductVariant,
    Quote,
    RecordSource,
    Vendor,
)

router = APIRouter(prefix="/excel", tags=["Excel 匯入匯出"])

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HEADERS = [
    "活動名稱", "活動開始日", "活動結束日", "活動類型",
    "藝人", "品項", "商品名稱", "規格",
    "售價(台幣)", "製作量", "實售量", "剩餘量",
    "成本(原幣)", "幣別", "匯率", "通路", "是否完售", "備註",
]
REQUIRED = ["活動名稱", "活動開始日", "藝人", "品項", "商品名稱", "售價(台幣)", "製作量"]

EVENT_TYPE_MAP = {"場販": "onsite", "通販": "online", "混合": "mixed"}
CHANNEL_MAP = {"現場": "onsite", "通販": "online", "預購": "preorder"}


def _to_date(value) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value).strip(), "%Y-%m-%d").date()
    except ValueError:
        return None


def _to_int(value) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        return int(float(value))
    except (ValueError, TypeError):
        return None


def _to_decimal(value) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        return Decimal(str(value))
    except InvalidOperation:
        return None


@router.get("/template")
def download_template():
    """匯入範本：一張填資料、一張說明。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "匯入資料"
    ws.append(HEADERS)
    ws.append([
        "2025 夏日祭場販", "2025-07-20", "2025-07-21", "場販",
        "星野銀河", "壓克力立牌", "星野銀河 壓克力立牌", "",
        450, 300, 250, "", 32.5, "CNY", 4.4, "現場", "否", "範例列，匯入前請刪除",
    ])
    ws.append([
        "2025 夏日祭場販", "2025-07-20", "", "",
        "星野銀河", "T-shirt", "星野銀河 T-shirt", "M",
        800, 100, "", 20, 180, "TWD", "", "現場", "否", "範例列：用剩餘量反推實售",
    ])

    doc = wb.create_sheet("說明")
    lines = [
        "【匯入範本說明】",
        "",
        "必填欄位：活動名稱、活動開始日(YYYY-MM-DD)、藝人、品項、商品名稱、售價(台幣)、製作量",
        "實售量與剩餘量擇一填（都填以實售量為準；剩餘量會用「製作量−剩餘量」反推實售）",
        "",
        "選填欄位與預設值：",
        "  活動結束日（預設＝開始日）／活動類型：場販、通販、混合（預設場販）",
        "  規格：如 S/M/L，同商品多規格就填多列（活動+藝人+商品名相同即視為同一商品）",
        "  成本(原幣)＋幣別＋匯率：幣別預設 TWD、匯率預設 1，台幣成本＝原幣×匯率",
        "  通路：現場、通販、預購（預設現場）／是否完售：是、否",
        "",
        "規則：",
        "  藝人、品項、活動不存在會自動建立；商品重複（同活動同藝人同名）會擋下",
        "  全部列驗證通過才會匯入，任何一列有錯就整批不寫（錯誤會逐列列出）",
        "  範例列請刪除後再匯入",
    ]
    for line in lines:
        doc.append([line])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type=XLSX,
        headers={"Content-Disposition": "attachment; filename=import_template.xlsx"},
    )


@router.post("/import", response_model=schemas.ImportReport)
async def import_excel(
    file: UploadFile = File(...),
    dry_run: bool = True,
    db: Session = Depends(get_db),
    cid: int = Depends(get_company_id),
):
    try:
        wb = load_workbook(BytesIO(await file.read()), data_only=True)
    except Exception:
        raise HTTPException(422, "無法讀取檔案，請確認是 .xlsx 格式")
    ws = wb["匯入資料"] if "匯入資料" in wb.sheetnames else wb.active

    header_cells = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    col = {str(h).strip(): i for i, h in enumerate(header_cells) if h}
    missing = [h for h in REQUIRED if h not in col]
    if missing:
        raise HTTPException(422, f"範本缺少必填欄位：{'、'.join(missing)}")

    # 既有資料快取（依名稱比對）
    artists = {a.name: a for a in db.scalars(
        select(Artist).where(Artist.company_id == cid)).all()}
    item_types = {t.name: t for t in db.scalars(
        select(ItemType).where(ItemType.company_id == cid)).all()}
    events = {e.name: e for e in db.scalars(
        select(Event).where(Event.company_id == cid)).all()}
    existing_products = {
        (p.event_id, p.artist_id, p.name)
        for p in db.scalars(select(Product).where(Product.company_id == cid)).all()
    }

    errors: list[schemas.ImportRowError] = []
    parsed = []  # 驗證通過的列
    new_artists: set[str] = set()
    new_item_types: set[str] = set()
    new_events: dict[str, dict] = {}
    seen_products: dict[tuple, dict] = {}  # 檔內的商品彙整（含既有活動的新商品）

    def err(row_no: int, message: str):
        errors.append(schemas.ImportRowError(row=row_no, message=message))

    total = 0
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        total += 1

        def get(name):
            i = col.get(name)
            return row[i] if i is not None and i < len(row) else None

        text = lambda name: (str(get(name)).strip() if get(name) is not None else "")

        event_name = text("活動名稱")
        artist_name = text("藝人")
        type_name = text("品項")
        product_name = text("商品名稱")
        if not (event_name and artist_name and type_name and product_name):
            err(row_no, "活動名稱/藝人/品項/商品名稱不可空白")
            continue

        start = _to_date(get("活動開始日"))
        if start is None:
            err(row_no, "活動開始日格式錯誤（要 YYYY-MM-DD）")
            continue
        end = _to_date(get("活動結束日")) or start
        if end < start:
            err(row_no, "活動結束日早於開始日")
            continue

        price = _to_decimal(get("售價(台幣)"))
        if price is None or price < 0:
            err(row_no, "售價(台幣)必須是 ≥ 0 的數字")
            continue
        production = _to_int(get("製作量"))
        if production is None or production <= 0:
            err(row_no, "製作量必須是正整數")
            continue

        sold = _to_int(get("實售量"))
        leftover = _to_int(get("剩餘量"))
        if sold is None and leftover is None:
            err(row_no, "實售量與剩餘量至少填一個")
            continue
        if sold is None:
            sold = production - leftover
        if sold < 0 or sold > production:
            err(row_no, f"實售量 {sold} 不合理（製作量 {production}）")
            continue

        currency = text("幣別").upper() or "TWD"
        rate = _to_decimal(get("匯率")) or (Decimal("1") if currency == "TWD" else None)
        if rate is None:
            err(row_no, f"幣別 {currency} 需要填匯率")
            continue
        cost = _to_decimal(get("成本(原幣)")) or Decimal("0")

        event_type_text = text("活動類型") or "場販"
        if event_type_text not in EVENT_TYPE_MAP:
            err(row_no, "活動類型只能是：場販、通販、混合")
            continue
        channel_text = text("通路") or "現場"
        if channel_text not in CHANNEL_MAP:
            err(row_no, "通路只能是：現場、通販、預購")
            continue
        sold_out_text = text("是否完售")
        if sold_out_text not in ("", "是", "否"):
            err(row_no, "是否完售只能填：是、否")
            continue

        # 活動：既有的沿用；新的以第一次出現的資訊為準
        if event_name not in events and event_name not in new_events:
            new_events[event_name] = {
                "start": start, "end": end,
                "event_type": EVENT_TYPE_MAP[event_type_text],
            }
        if artist_name not in artists:
            new_artists.add(artist_name)
        if type_name not in item_types:
            new_item_types.add(type_name)

        # 商品彙整：同活動×藝人×商品名 = 同商品的多規格
        pkey = (event_name, artist_name, product_name)
        existing_event = events.get(event_name)
        if (
            existing_event is not None
            and artist_name in artists
            and (existing_event.id, artists[artist_name].id, product_name)
            in existing_products
        ):
            err(row_no, f"商品「{product_name}」已存在於活動「{event_name}」（疑似重複匯入）")
            continue
        if pkey not in seen_products:
            seen_products[pkey] = {"price": price, "type": type_name, "variants": {}}
        variant_name = text("規格") or "單一規格"
        if variant_name in seen_products[pkey]["variants"]:
            err(row_no, f"商品「{product_name}」的規格「{variant_name}」重複出現")
            continue
        seen_products[pkey]["variants"][variant_name] = True

        parsed.append({
            "event": event_name, "artist": artist_name, "type": type_name,
            "product": product_name, "variant": variant_name,
            "price": price, "production": production, "sold": sold,
            "cost": cost, "currency": currency, "rate": rate,
            "channel": CHANNEL_MAP[channel_text],
            "sold_out": sold_out_text == "是",
            "notes": text("備註") or None,
        })

    report = schemas.ImportReport(
        total=total, valid=len(parsed), errors=errors,
        new_events=len(new_events), new_artists=len(new_artists),
        new_item_types=len(new_item_types), new_products=len(seen_products),
        imported=False,
    )
    if dry_run or errors:
        return report

    # ── 正式寫入（全部驗證通過才會走到這）──
    for name in new_artists:
        artists[name] = Artist(company_id=cid, name=name)
        db.add(artists[name])
    for name in new_item_types:
        item_types[name] = ItemType(company_id=cid, name=name)
        db.add(item_types[name])
    for name, info in new_events.items():
        events[name] = Event(
            company_id=cid, name=name, start_date=info["start"],
            end_date=info["end"], event_type=info["event_type"],
        )
        db.add(events[name])
    db.flush()

    products: dict[tuple, Product] = {}
    for row_data in parsed:
        pkey = (row_data["event"], row_data["artist"], row_data["product"])
        if pkey not in products:
            products[pkey] = Product(
                company_id=cid,
                event_id=events[row_data["event"]].id,
                artist_id=artists[row_data["artist"]].id,
                item_type_id=item_types[row_data["type"]].id,
                name=row_data["product"],
                price_twd=row_data["price"],
            )
            db.add(products[pkey])
            db.flush()
        product = products[pkey]
        event = events[row_data["event"]]

        variant = ProductVariant(
            company_id=cid,
            product_id=product.id,
            variant_name=row_data["variant"],
            production_qty=row_data["production"],
            cost_amount=row_data["cost"],
            cost_currency=row_data["currency"],
            exchange_rate=row_data["rate"],
            cost_twd=(row_data["cost"] * row_data["rate"]).quantize(Decimal("0.01")),
            notes=row_data["notes"],
        )
        db.add(variant)
        db.flush()

        db.add(InventoryMovement(
            company_id=cid, variant_id=variant.id,
            movement_type=MovementType.INBOUND.value,
            quantity_delta=row_data["production"],
            movement_date=event.start_date,
            source=RecordSource.EXCEL.value,
        ))
        if row_data["sold"] > 0:
            db.add(InventoryMovement(
                company_id=cid, variant_id=variant.id,
                movement_type=MovementType.SALE.value,
                quantity_delta=-row_data["sold"],
                movement_date=event.end_date,
                channel=row_data["channel"],
                sale_price_twd=row_data["price"],
                sold_out_today=row_data["sold_out"],
                source=RecordSource.EXCEL.value,
            ))
    db.commit()
    report.imported = True
    return report


@router.get("/export")
def export_excel(db: Session = Depends(get_db), cid: int = Depends(get_company_id)):
    """把公司全部業務資料匯出成一個 Excel（備份/離線分析用）。"""
    wb = Workbook()

    ws = wb.active
    ws.title = "商品銷售明細"
    ws.append(["活動", "開始日", "藝人", "品項", "商品", "規格", "售價(台幣)",
               "成本(台幣)", "製作量", "實售量", "目前庫存", "營收(台幣)"])
    products = db.scalars(
        select(Product).where(Product.company_id == cid)
        .options(selectinload(Product.variants), selectinload(Product.artist),
                 selectinload(Product.event), selectinload(Product.item_type))
        .order_by(Product.event_id, Product.artist_id, Product.id)
    ).all()
    movements = db.scalars(
        select(InventoryMovement).where(InventoryMovement.company_id == cid)
    ).all()
    stock: dict[int, int] = {}
    sold: dict[int, int] = {}
    revenue: dict[int, Decimal] = {}
    for m in movements:
        stock[m.variant_id] = stock.get(m.variant_id, 0) + m.quantity_delta
        if m.movement_type in ("sale", "sale_return"):
            sold[m.variant_id] = sold.get(m.variant_id, 0) - m.quantity_delta
            revenue[m.variant_id] = revenue.get(m.variant_id, Decimal(0)) + (
                -m.quantity_delta * (m.sale_price_twd or Decimal(0))
            )
    for p in products:
        for v in p.variants:
            ws.append([
                p.event.name, str(p.event.start_date), p.artist.name,
                p.item_type.name if p.item_type else "套組",
                p.name, v.variant_name, float(p.price_twd), float(v.cost_twd),
                v.production_qty, sold.get(v.id, 0), stock.get(v.id, 0),
                float(revenue.get(v.id, 0)),
            ])

    ws2 = wb.create_sheet("報價紀錄")
    ws2.append(["日期", "品項", "廠商", "數量", "單價(原幣)", "幣別", "匯率", "單價(台幣)", "備註"])
    quotes = db.execute(
        select(Quote, ItemType.name, Vendor.name)
        .join(ItemType, Quote.item_type_id == ItemType.id)
        .outerjoin(Vendor, Quote.vendor_id == Vendor.id)
        .where(Quote.company_id == cid).order_by(Quote.quote_date)
    ).all()
    for q, tname, vname in quotes:
        ws2.append([str(q.quote_date), tname, vname or "", q.quantity,
                    float(q.unit_price), q.currency, float(q.exchange_rate),
                    float(q.unit_price_twd), q.notes or ""])

    ws3 = wb.create_sheet("檔期開支")
    ws3.append(["活動", "類別", "金額(台幣)", "歸屬藝人", "備註"])
    cat = {"booth": "攤位費", "shipping": "運費", "labor": "人力", "other": "其他"}
    expenses = db.execute(
        select(Expense, Event.name, Artist.name)
        .join(Event, Expense.event_id == Event.id)
        .outerjoin(Artist, Expense.artist_id == Artist.id)
        .where(Expense.company_id == cid).order_by(Expense.event_id)
    ).all()
    for e, ename, aname in expenses:
        ws3.append([ename, cat.get(e.category, e.category), float(e.amount_twd),
                    aname or "", e.notes or ""])

    ws4 = wb.create_sheet("藝人熱度")
    ws4.append(["藝人", "日期", "平台", "指標", "數值"])
    metrics = db.execute(
        select(ArtistMetric, Artist.name)
        .join(Artist, ArtistMetric.artist_id == Artist.id)
        .where(ArtistMetric.company_id == cid)
        .order_by(ArtistMetric.artist_id, ArtistMetric.record_date)
    ).all()
    for m, aname in metrics:
        ws4.append([aname, str(m.record_date), m.platform, m.metric_type, m.value])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type=XLSX,
        headers={"Content-Disposition": "attachment; filename=vtuber_merch_export.xlsx"},
    )
